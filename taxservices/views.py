from io import BytesIO

from django.db.models import Q
from django.http import HttpResponse, HttpResponseRedirect
from django.shortcuts import render, get_object_or_404
from django.views.generic import View
from django_pandas.io import read_frame
from django_tables2.export import TableExport
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Border, Side

from credits.decorators import group_required
from taxservices.models import *
from taxservices.tables import *
from docxtpl import DocxTemplate

from .forms import *

class Search(View):
    def get(self, request):
        return render(request, 'taxservices/search_form.html')

    def post(self, request):
        bound_form = SearchForm(request.POST)
        if bound_form.is_valid():

            if bound_form.cleaned_data['type'] == 'J':
                model = JuridicalPerson.objects.filter(tin=bound_form.cleaned_data['tin']).first()
                if (not model):
                    model = JuridicalPerson.services.get('yurnp1', {
                        "company_tin": bound_form.cleaned_data['tin'],
                        "lang":"uz"})
            else:
                model = PhysicalPerson.objects.filter(tin=bound_form.cleaned_data['tin']).first()
                if (not model):
                    model = PhysicalPerson.services.get('fiznp1', {
                        "tin": bound_form.cleaned_data['tin'],
                        "pinfl": bound_form.cleaned_data['pinfl'],
                        "series_passport": bound_form.cleaned_data['series'],
                        "number_passport": bound_form.cleaned_data['number'],
                        "lang": "uz"})
            if model is not None:
                model.save()
                return HttpResponseRedirect(reverse('payer_detail_url', args=(model.slug,)))
            else:
                bound_form.add_error('tin','Информации об этом ИНН не найдено.')

        return render(request, 'taxservices/search_form.html', context={'form': bound_form})


# JURIDICAL PERSON SERVICES
class PayerDetail(View):
    def get(self, request, slug):
        if slug[0] == 'J':
            person = JuridicalPerson.objects.filter(slug=slug).first()
            return render(request, 'taxservices/payer_detail_juridical.html', context={'person': person})
        else:
            person = PhysicalPerson.objects.get(slug=slug)
            return render(request, 'taxservices/payer_detail_physical.html', context={'person': person})

class EmployeeList(View):
    def get(self, request, slug):
        person = JuridicalPerson.objects.filter(slug=slug).first()
        employee = Employee.services.all('yur-employee-count', {"tin": person.tin, "year": "2019", "lang": "uz"})
        table = EmployeeListTable(employee)
        context = {
            'person': person,
            'table': table
        }
        return render(request, 'taxservices/employee_list.html', context=context)

class BalancesList(View):
    def get(self, request, slug):
        person = JuridicalPerson.objects.filter(slug=slug).first()
        balances = Balance.services.all('buxbalans', {"tin": person.tin, "year": "2019", "period": "4", "lang": "uz"})
        table = BalancesListTable(balances)
        context = {
            'person': person,
            'table': table
        }
        return render(request, 'taxservices/balances_list.html', context=context)

class FinancesList(View):
    def get(self, request, slug):
        person = JuridicalPerson.objects.filter(slug=slug).first()
        finances = Finance.services.all('finreport', {"tin": person.tin, "year": "2019", "period": "4", "lang": "uz"})
        table = FinancesListTable(finances)
        context = {
            'person': person,
            'table': table
        }
        return render(request, 'taxservices/finances_list.html', context=context)


class ViewBaseNDS(View):
    def get(self, request, slug):
        person = JuridicalPerson.objects.filter(slug=slug).first()
        base_nds = BaseNDS.services.get('nds', {"tin": person.tin, "year": "2019", "period": "4", "lang": "uz"})
        context = {
            'person': person,
            'object': base_nds
        }
        return render(request, 'taxservices/base_nds_info.html', context=context)

class ViewReportENP(View):
    def get(self, request, slug):
        person = JuridicalPerson.objects.filter(slug=slug).first()
        report_enp = ReportENP.services.get('enp', {"tin": person.tin, "year": "2019", "period": "4", "lang": "uz"})
        context = {
            'person': person,
            'object': report_enp
        }
        return render(request, 'taxservices/report_enp_info.html', context=context)


# PHYSICAL PERSON SERVICES
class SalariesList(View):
    def get(self, request, slug):
        person = PhysicalPerson.objects.filter(slug=slug).first()
        salaries = Salary.services.all('fiz-salary', {
            "tin": person.tin, "pinfl": "", "lang": "uz",
            "series_passport": "", "number_passport": ""})
        table = SalaryTable(salaries)
        context = {
            'title': 'Сведения о заработной плате',
            'person': person,
            'table': table
        }
        return render(request, 'taxservices/includes/list.physic.html', context=context)


class TaxDebtsList(View):
    def get(self, request, slug):
        if slug[0] == 'J':
            person = JuridicalPerson.objects.get(slug=slug)
            taxdebts = TaxDebt.services.all('yurdebt', {"tin": person.tin, "lang": "uz"})
            table = TaxDebtTable(taxdebts)
            context = {
                'person': person,
                'table': table
            }
            return render(request, 'taxservices/taxdebts_list.html', context=context)
        else:
            person = PhysicalPerson.objects.get(slug=slug)
            taxdebts = TaxDebt.services.all('fizdebt', {"tin": person.tin, "lang": "uz"})
            table = TaxDebtTable(taxdebts)
            context = {
                'title': 'Налоговая задолженность',
                'person': person,
                'table': table
            }
            return render(request, 'taxservices/includes/list.physic.html', context=context)


class FoundersList(View):
    def get(self, request, slug):
        person = PhysicalPerson.objects.filter(slug=slug).first()
        founders = Founder.services.all('np3', {"tin": person.tin, "lang": "uz"})
        table = FounderTable(founders)
        context = {
            'title': 'Юридическое лицо, в котором физическое лицо является учредителем',
            'person': person,
            'table': table
        }
        return render(request, 'taxservices/includes/list.physic.html', context=context)


class DividendList(View):
    def get(self, request, slug):
        person = PhysicalPerson.objects.filter(slug=slug).first()
        dividend = Dividend.services.all('np3', {"tin": person.tin, "year": "2019", "lang": "uz"})
        table = DividendTable(dividend)
        context = {
            'title': 'Сведения о полученных дивидендах',
            'person': person,
            'table': table
        }
        return render(request, 'taxservices/includes/list.physic.html', context=context)


class ObjectsList(View):
    def get(self, request, slug):
        person = PhysicalPerson.objects.filter(slug=slug).first()
        objects = Object.services.all('fiz/tax-objects', {"tin": person.tin, "lang": "uz"})
        table = ObjectTable(objects)
        context = {
            'title': 'Сведения об имуществе физического лица',
            'person': person,
            'table': table
        }
        return render(request, 'taxservices/includes/list.physic.html', context=context)


class LeasedsList(View):
    def get(self, request, slug):
        person = PhysicalPerson.objects.filter(slug=slug).first()
        leaseds = Leased.services.all('rents', {"tin": person.tin, "lang": "uz"})
        table = LeasedTable(leaseds)
        context = {
            'title': 'Сведения о сданной в аренду имущества',
            'person': person,
            'table': table
        }
        return render(request, 'taxservices/includes/list.physic.html', context=context)


class GenerateExcel(View):
    def get(self, request, slug):
        person = JuridicalPerson.objects.filter(slug=slug).first()

        # Налоговый задолженность
        taxdebts = TaxDebt.services.all('yurdebt', {"tin": person.tin, "lang": "uz"})
        mydict1 = []
        for item in taxdebts:
            mydict1.append({
                'tax_name': item.get_tax_name(),
                'ned_date': item.ned_date,
                'ned_summa': item.ned_summa,
                'pen_summa': item.pen_summa,
            })
        df_taxdebt = pd.DataFrame(mydict1)

        # Количество сотрудника
        employee = Employee.services.all('yur-employee-count', {"tin": person.tin, "year": "2019", "lang": "uz"})
        mydict = []
        for item in employee:
            mydict.append({
                'fio': item.fio,
                'tin': item.tin,
                'doxod1': item.doxod1,
                'doxod2': item.doxod2,
            })
        df_employee = pd.DataFrame(mydict)

        # Форма 1
        balances = Balance.services.dictall('buxbalans', {
            "tin": person.tin, "year": "2019",
            "period": "4", "lang": "uz"})
        df_right = pd.DataFrame(list(balances.values()))
        df_balance = read_frame(BalanceSheet.objects.all())[['code_str']]
        df_balance = df_balance.join(df_right.set_index('row_no'), on='code_str')

        # Форма 2
        finances = Finance.services.dictall('finreport', {
            "tin": person.tin, "year": "2019",
            "period": "4", "lang": "uz"})
        df_right = pd.DataFrame(list(finances.values()))
        df_finance = read_frame(FinanceReport.objects.all())[['code_str']]
        df_finance = df_finance.join(df_right.set_index('row_no'), on='code_str')

        # ЕНП
        report_enp = ReportENP.services.get('enp', {"tin": person.tin, "year": "2019", "period": "4", "lang": "uz"})
        base_nds = BaseNDS.services.get('nds', {"tin": person.tin, "year": "2019", "period": "4", "lang": "uz"})

        with BytesIO() as b:

            workbook = load_workbook('media/tmp_files/modul.xlsx')
            writer = pd.ExcelWriter(b, engine='openpyxl')
            writer.book = workbook
            writer.sheets = dict((ws.title, ws) for ws in workbook.worksheets)

            df_taxdebt.to_excel(writer, sheet_name='Налоговая задолженность', startrow=2, startcol=0, header=False, index=False)
            df_employee.to_excel(writer, sheet_name='Количество работников', startrow=2, startcol=0, header=False, index=False)

            df_balance.to_excel(writer, sheet_name='Форма № 1', startrow=3, startcol=1, header=False, index=False)
            df_finance.to_excel(writer, sheet_name='Форма № 2', startrow=4, startcol=1, header=False, index=False)

            if report_enp:
                sheet = workbook["ЕНП"]
                sheet['B2'].value = report_enp.enp101
                sheet['B3'].value = report_enp.enp102
                sheet['B4'].value = report_enp.prib10

            if base_nds:
                sheet = workbook["НДС"]
                sheet['B2'].value = base_nds.nds
                sheet['B3'].value = base_nds.ndsupr

            writer.save()
            response = HttpResponse(b.getvalue(), content_type='application/vnd.ms-excel')
            response["Content-Disposition"] = 'attachment; filename="tax_report.xlsx"'
            return response

class GenerateWord(View):
    def get(self, request, slug):
        person = PhysicalPerson.objects.filter(slug=slug).first()
        salaries = Salary.services.all('fiz-salary', {
            "tin": person.tin, "pinfl": "", "lang": "uz",
            "series_passport": "", "number_passport": ""})
        taxdebts = TaxDebt.services.all('fizdebt', {"tin": person.tin, "lang": "uz"})
        founders = Founder.services.all('np3', {"tin": person.tin, "lang": "uz"})
        dividend = Dividend.services.all('np3', {"tin": person.tin, "year": "2019", "lang": "uz"})
        objects = Object.services.all('fiz/tax-objects', {"tin": person.tin, "lang": "uz"})
        leaseds = Leased.services.all('rents', {"tin": person.tin, "lang": "uz"})

        doc = DocxTemplate("media/tmp_files/report.docx")
        context = {
            'person': person,
            'salaries': salaries,
            'taxdebts': taxdebts,
            'founders': founders,
            'dividend': dividend,
            'objects': objects,
            'leaseds': leaseds
        }

        doc.render(context)

        byte_io = BytesIO()
        doc.save(byte_io)
        byte_io.seek(0)

        response = HttpResponse(byte_io.read())

        # Content-Disposition header makes a file downloadable
        response["Content-Disposition"] = "attachment; filename=generated_doc.docx"

        # Set the appropriate Content-Type for docx file
        response["Content-Type"] = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        return response
